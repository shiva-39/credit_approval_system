"""
Celery tasks for background data ingestion.
"""
from celery import shared_task
from django.db import transaction
import openpyxl
from datetime import datetime
from decimal import Decimal
import os
from .models import Customer


@shared_task
def ingest_customer_data():
    """
    Ingest customer data from customer_data.xlsx file.
    
    Expected columns:
    - customer_id
    - first_name
    - last_name
    - phone_number
    - monthly_salary
    - approved_limit
    - current_debt
    """
    file_path = 'customer_data.xlsx'
    
    if not os.path.exists(file_path):
        return f"File {file_path} not found"
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        customers_created = 0
        customers_updated = 0
        
        # Skip header row
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Skip empty rows
                continue
            
            customer_id = int(row[0]) if row[0] else None
            first_name = str(row[1]) if row[1] else ''
            last_name = str(row[2]) if row[2] else ''
            phone_number = int(row[3]) if row[3] else 0
            monthly_salary = Decimal(str(row[4])) if row[4] else Decimal('0')
            approved_limit = Decimal(str(row[5])) if row[5] else Decimal('0')
            current_debt = Decimal(str(row[6])) if row[6] else Decimal('0')
            
            with transaction.atomic():
                customer, created = Customer.objects.update_or_create(
                    customer_id=customer_id,
                    defaults={
                        'first_name': first_name,
                        'last_name': last_name,
                        'phone_number': phone_number,
                        'monthly_salary': monthly_salary,
                        'approved_limit': approved_limit,
                        'current_debt': current_debt,
                    }
                )
                
                if created:
                    customers_created += 1
                else:
                    customers_updated += 1
        
        workbook.close()
        
        return f"Customer data ingestion completed. Created: {customers_created}, Updated: {customers_updated}"
    
    except Exception as e:
        return f"Error ingesting customer data: {str(e)}"


@shared_task
def ingest_loan_data():
    """
    Ingest loan data from loan_data.xlsx file.
    
    Expected columns:
    - customer_id
    - loan_id
    - loan_amount
    - tenure
    - interest_rate
    - monthly_repayment
    - emis_paid_on_time
    - start_date
    - end_date
    """
    from loans.models import Loan
    
    file_path = 'loan_data.xlsx'
    
    if not os.path.exists(file_path):
        return f"File {file_path} not found"
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        loans_created = 0
        loans_updated = 0
        
        # Skip header row
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:  # Skip empty rows
                continue
            
            customer_id = int(row[0]) if row[0] else None
            loan_id = int(row[1]) if row[1] else None
            loan_amount = Decimal(str(row[2])) if row[2] else Decimal('0')
            tenure = int(row[3]) if row[3] else 0
            interest_rate = Decimal(str(row[4])) if row[4] else Decimal('0')
            monthly_repayment = Decimal(str(row[5])) if row[5] else Decimal('0')
            emis_paid_on_time = int(row[6]) if row[6] else 0
            start_date = row[7] if row[7] else datetime.now().date()
            end_date = row[8] if row[8] else datetime.now().date()
            
            # Convert dates if they're strings
            if isinstance(start_date, str):
                start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
            if isinstance(end_date, str):
                end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
            
            try:
                customer = Customer.objects.get(customer_id=customer_id)
            except Customer.DoesNotExist:
                continue  # Skip loans for non-existent customers
            
            with transaction.atomic():
                loan, created = Loan.objects.update_or_create(
                    loan_id=loan_id,
                    defaults={
                        'customer': customer,
                        'loan_amount': loan_amount,
                        'tenure': tenure,
                        'interest_rate': interest_rate,
                        'monthly_repayment': monthly_repayment,
                        'emis_paid_on_time': emis_paid_on_time,
                        'start_date': start_date,
                        'end_date': end_date,
                    }
                )
                
                if created:
                    loans_created += 1
                else:
                    loans_updated += 1
        
        workbook.close()
        
        return f"Loan data ingestion completed. Created: {loans_created}, Updated: {loans_updated}"
    
    except Exception as e:
        return f"Error ingesting loan data: {str(e)}"


@shared_task
def ingest_all_data():
    """
    Ingest both customer and loan data.
    """
    customer_result = ingest_customer_data()
    loan_result = ingest_loan_data()
    
    return f"{customer_result}\n{loan_result}"
